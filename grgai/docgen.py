"""docgen.py — turn a JSON spec into a real .pptx / .xlsx / .docx / .pdf file (bytes)."""
import io
import os
import re

from pptx import Presentation
import openpyxl
from openpyxl.styles import Font, PatternFill
from docx import Document

MIME = {
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
MAX_SLIDES, MAX_ROWS, MAX_SHEETS = 120, 8000, 25

# Bundled Unicode font so PDFs render Romanian diacritics (ș ț ă â î) correctly.
_FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
_FONT_REG = os.path.join(_FONT_DIR, "DejaVuSans.ttf")
_FONT_BOLD = os.path.join(_FONT_DIR, "DejaVuSans-Bold.ttf")


def _safe_name(name, ext):
    n = re.sub(r"[^A-Za-z0-9 ._-]", "", str(name or "document")).strip() or "document"
    if not n.lower().endswith("." + ext):
        n = re.sub(r"\.(pptx|xlsx|docx|pdf)$", "", n, flags=re.I) + "." + ext
    return n[:80]


def _pptx(spec):
    p = Presentation()
    if spec.get("title"):
        s = p.slides.add_slide(p.slide_layouts[0])
        s.shapes.title.text = str(spec["title"])[:200]
        if spec.get("subtitle") and len(s.placeholders) > 1:
            s.placeholders[1].text = str(spec["subtitle"])[:300]
    for sl in (spec.get("slides") or [])[:MAX_SLIDES]:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = str(sl.get("title") or "")[:200]
        if len(s.placeholders) > 1:
            body = s.placeholders[1].text_frame
            bullets = sl.get("bullets") or ([sl.get("content")] if sl.get("content") else [])
            first = True
            for b in bullets[:40]:
                if first:
                    body.text = str(b)[:500]; first = False
                else:
                    body.add_paragraph().text = str(b)[:500]
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = str(sl["notes"])[:2000]
    buf = io.BytesIO(); p.save(buf); return buf.getvalue()


def _xlsx(spec):
    wb = openpyxl.Workbook()
    sheets = spec.get("sheets") or [{"name": spec.get("title") or "Sheet1",
                                     "headers": spec.get("headers") or [], "rows": spec.get("rows") or []}]
    first = True
    for sh in sheets[:MAX_SHEETS]:
        ws = wb.active if first else wb.create_sheet(); first = False
        ws.title = (str(sh.get("name") or "Sheet")[:31]) or "Sheet"
        headers = sh.get("headers") or []
        if headers:
            ws.append([str(h) for h in headers])
            for c in ws[1]:
                c.font = Font(bold=True, color="4C1D95")
                c.fill = PatternFill("solid", fgColor="EDE9FE")
        for row in (sh.get("rows") or [])[:MAX_ROWS]:
            ws.append(list(row) if isinstance(row, (list, tuple)) else [row])
        for col in ws.columns:
            try:
                w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, w + 2))
            except Exception:
                pass
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _docx(spec):
    d = Document()
    if spec.get("title"):
        d.add_heading(str(spec["title"])[:200], 0)
    secs = spec.get("sections") or []
    for sec in secs:
        if isinstance(sec, str):          # tolerate a section given as plain text
            if sec.strip():
                d.add_paragraph(sec)
            continue
        if not isinstance(sec, dict):
            continue
        if sec.get("heading"):
            lvl = int(sec.get("level", 1) or 1)
            d.add_heading(str(sec["heading"])[:200], max(1, min(4, lvl)))
        for para in (sec.get("paragraphs") or []):
            d.add_paragraph(str(para))
        for b in (sec.get("bullets") or []):
            d.add_paragraph(str(b), style="List Bullet")
    if not secs:                          # tolerate a flat docx (top-level paragraphs/bullets/body)
        for para in (spec.get("paragraphs") or []):
            d.add_paragraph(str(para))
        for b in (spec.get("bullets") or []):
            d.add_paragraph(str(b), style="List Bullet")
        if spec.get("body"):
            for para in str(spec["body"]).split("\n\n"):
                if para.strip():
                    d.add_paragraph(para.strip())
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def _pdf(spec):
    from fpdf import FPDF  # lazy import: docx/xlsx/pptx keep working even if fpdf2 is absent
    try:
        from fpdf.enums import XPos, YPos
        _NX, _NY = XPos.LMARGIN, YPos.NEXT
    except Exception:  # noqa: BLE001
        _NX = _NY = None

    have_reg = os.path.exists(_FONT_REG)
    have_bold = os.path.exists(_FONT_BOLD)
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(20, 18, 20)
    pdf.add_page()
    if have_reg:
        pdf.add_font("Doc", "", _FONT_REG)
        if have_bold:
            pdf.add_font("Doc", "B", _FONT_BOLD)
        fam = "Doc"
    else:
        fam = "Helvetica"
    bold = "B" if (have_bold or fam == "Helvetica") else ""

    def _t(s):
        s = str(s)
        return s if have_reg else s.encode("latin-1", "replace").decode("latin-1")

    def _line(text, size=11, style="", h=6.2, align="J", gap=2.0):
        pdf.set_font(fam, style, size)
        if _NX is not None:
            pdf.multi_cell(0, h, _t(text), align=align, new_x=_NX, new_y=_NY)
        else:
            pdf.multi_cell(0, h, _t(text), align=align)
        if gap:
            pdf.ln(gap)

    if spec.get("title"):
        _line(spec["title"], size=15, style=bold, h=8, align="C", gap=4)

    def _block(paras, bullets, heading=None, level=1):
        if heading:
            _line(heading, size=(13 if level <= 1 else 12), style=bold, h=7, align="L", gap=1.5)
        for p in (paras or []):
            _line(p, size=11, style="", h=6.2, align="J", gap=2)
        for b in (bullets or []):
            _line("•  " + str(b), size=11, style="", h=6.2, align="L", gap=1)

    secs = spec.get("sections") or []
    for sec in secs:
        if isinstance(sec, str):
            if sec.strip():
                _block([sec], None)
            continue
        if not isinstance(sec, dict):
            continue
        _block(sec.get("paragraphs"), sec.get("bullets"), sec.get("heading"), int(sec.get("level", 1) or 1))
    if not secs:
        _block(spec.get("paragraphs"), spec.get("bullets"))
        if spec.get("body"):
            for para in str(spec["body"]).split("\n\n"):
                if para.strip():
                    _block([para.strip()], None)

    return bytes(pdf.output())


def generate(spec):
    """spec: {type: 'pptx'|'xlsx'|'docx'|'pdf', filename, title, ...}. Returns (bytes, mime, filename)."""
    if not isinstance(spec, dict):
        raise ValueError("spec must be an object")
    t = (spec.get("type") or "").lower()
    if t not in MIME:                     # infer from shape so a missing/typo'd type still builds
        if spec.get("slides"):
            t = "pptx"
        elif spec.get("sheets") or spec.get("rows") or spec.get("headers"):
            t = "xlsx"
        elif spec.get("sections") or spec.get("paragraphs") or spec.get("body"):
            t = "docx"
    if t not in MIME:
        raise ValueError("type must be pptx, xlsx, docx or pdf")
    data = {"pptx": _pptx, "xlsx": _xlsx, "docx": _docx, "pdf": _pdf}[t](spec)
    return data, MIME[t], _safe_name(spec.get("filename") or spec.get("title") or "document", t)
